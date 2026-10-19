"""Importa a aba "Lading Page" da planilha antiga para a lista "Landing Pages" do crm.xlsx.

Uso:  python importar.py "C:\\caminho\\Controle de Lading Page.xlsx"

Se a lista "Landing Pages" já existir no crm.xlsx, o script para sem mexer em nada. Para substituí-la
de propósito (perdendo o que foi feito no CRM), acrescente --substituir (o backup é feito antes).
"""
import datetime as dt
import sys

import openpyxl

import server

ABA = "Lading Page"
LISTA = "Landing Pages"
IGNORAR = {"Conta", "SLA"}  # SLA é fórmula; "Dias na etapa" substitui
TIPOS = {"Ferramenta": "selecao", "Segmento": "selecao", "Origem": "selecao", "Motivo": "selecao",
         "Endereço": "url", "Data inicio": "data", "Data fim": "data"}
RENOMEAR = {"Data inicio": "Data início"}
# (nome, tipo, cor, oferta) — "oferta": etapas depois da proposta; o funil calcula o % delas sobre o total de ofertas
ETAPAS = [("Baixar fotos", "aberta", "#94a3b8", False), ("Desenvolver LP", "aberta", "#60a5fa", False),
          ("Entrar em contato", "aberta", "#38bdf8", False), ("Visita inicial", "aberta", "#a78bfa", False),
          ("Follow", "aberta", "#fbbf24", True), ("Fechamento", "aberta", "#fb923c", True),
          ("Sem resposta", "aberta", "#94a3b8", True), ("Venda", "ganha", "#22c55e", True),
          ("Recusa", "perdida", "#ef4444", True), ("Desqualificado", "perdida", "#6b7280", False)]


def importar(origem, substituir=False):
    ws = openpyxl.load_workbook(origem, data_only=True)[ABA]
    cabecalho = []
    for cell in ws[1]:
        if cell.value is None:  # depois da 1ª coluna vazia começa a área do Funil/Projeção
            break
        cabecalho.append(str(cell.value).strip())

    colunas, mapa = [], {}  # mapa: índice na planilha -> nome da coluna no CRM
    for i, h in enumerate(cabecalho):
        if h in IGNORAR:
            continue
        nome, tipo = ("Etapa", "etapa") if h == "Status" else (RENOMEAR.get(h, h), TIPOS.get(h, "texto"))
        colunas.append(server.col(nome, tipo))
        mapa[i] = nome
    colunas += [server.col(n, t) for n, t in server.COLUNAS_SISTEMA if t != "etapa"]
    tipos = {c["nome"]: c["tipo"] for c in colunas}

    registros = []
    for linha in ws.iter_rows(min_row=2, max_col=len(cabecalho), values_only=True):
        if not linha[0]:
            continue
        r = {"_id": server.novo_id()}
        for i, nome in mapa.items():
            v = server.ler_valor(linha[i], tipos[nome])
            if v is not None:
                r[nome] = v
        datas = [r.get("Data fim"), r.get("Data início")]
        r["Etapa desde"] = next((d for d in datas if d and len(d) == 10 and d[4] == "-"), dt.date.today().isoformat())
        registros.append(r)

    for c in colunas:  # opções = valores encontrados na planilha
        if c["tipo"] == "selecao":
            c["opcoes"] = sorted({r[c["nome"]] for r in registros if c["nome"] in r}, key=str.casefold)

    lista = {"nome": LISTA, "colunas": colunas,
             "etapas": [{"nome": n, "tipo": t, "cor": cor, "oferta": o} for n, t, cor, o in ETAPAS], "registros": registros}
    server.completar(lista)  # etapa que não esteja na lista acima entra como "aberta"

    estado = server.ler() if server.ARQ.exists() else {"listas": []}
    if any(l["nome"] == LISTA for l in estado["listas"]) and not substituir:
        sys.exit(f'A lista "{LISTA}" já existe no crm.xlsx e seria SUBSTITUÍDA (perdendo o que foi feito no CRM).\n'
                 f'Se é isso mesmo, rode de novo com --substituir no final.')
    estado["listas"] = [lista] + [l for l in estado["listas"] if l["nome"] != LISTA]
    server.gravar(estado)
    print(f'{len(registros)} registros importados para a lista "{LISTA}" em {server.ARQ}')


if __name__ == "__main__":
    args = [a for a in sys.argv[1:] if a != "--substituir"]
    if len(args) != 1:
        sys.exit(__doc__)
    importar(args[0], substituir="--substituir" in sys.argv)
