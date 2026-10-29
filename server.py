"""CRM local: serve o index.html e lê/grava o crm.xlsx. Só escuta em 127.0.0.1."""
import datetime as dt
import json
import os
import re
import shutil
import sys
import time
import uuid
import webbrowser
from http.server import BaseHTTPRequestHandler, HTTPServer
from pathlib import Path

try:
    import openpyxl
    from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter
except ImportError:
    sys.exit("Falta o openpyxl. Rode no terminal:  python -m pip install openpyxl")

PASTA = Path(__file__).resolve().parent
ARQ = PASTA / "crm.xlsx"
BACKUPS = PASTA / "backups"
MAX_BACKUPS = 20
INTERVALO_BACKUP = 600  # segundos: no máximo um backup a cada 10 min (e sempre no 1º salvamento)
PORTA = 8765
URL = f"http://127.0.0.1:{PORTA}/"

ABA_CONFIG = "_Config"
CAB_CONFIG = ["Lista", "Item", "Nome", "Tipo", "Opções", "Ordem", "Visível", "Cor", "Oferta", "Produção", "Limite dias", "Padrão"]
ABA_HIST = "_Historico"  # mudanças de etapa: uma linha por mudança
CAB_HIST = ["Lista", "_id", "Data", "De", "Para"]
ABA_ITENS = "_Servicos"  # serviços vendidos: uma linha por item de cada registro
CAB_ITENS = ["Lista", "_id", "Serviço", "Cobrança", "Valor"]
ABA_CONTATOS = "_Contatos"  # contatos de contabilidades e clientes: uma linha por contato
CAB_CONTATOS = ["Tipo", "_id do registro", "Nome", "Cargo", "Telefone", "E-mail", "Observação", "_id"]
CAMPOS_CONTATO = ["nome", "cargo", "telefone", "email", "obs"]
VERSAO_API = 2  # o index.html só cria Contabilidades/Clientes/Contatos com esta versão ou mais nova
COBRANCAS = {"unico": "Único", "mensal": "Mensal"}
# "vinculo": guarda o _id de um registro de outra lista (ex.: o cliente aponta para a contabilidade)
TIPOS_COLUNA = ["texto", "numero", "data", "selecao", "url", "telefone", "caixa", "servicos", "vinculo"]
# Colunas que toda lista tem. O tipo identifica a coluna, então o nome pode ser trocado.
COLUNAS_SISTEMA = [("Etapa", "etapa"), ("Próxima ação", "proxima_acao"),
                   ("Data da próxima ação", "data_proxima_acao"),
                   ("Etapa desde", "etapa_desde"), ("Dias na etapa", "dias_na_etapa")]
TIPOS_SISTEMA = [t for _, t in COLUNAS_SISTEMA]
TIPOS_DATA = ["data", "etapa_desde", "data_proxima_acao"]
TIPOS_ETAPA = ["aberta", "ganha", "perdida"]

ultimo_backup = 0.0


class ArquivoTravado(Exception):
    pass


def col(nome, tipo, opcoes=None, visivel=True):
    return {"nome": nome, "tipo": tipo, "opcoes": opcoes or [], "visivel": visivel}


def nova_lista(nome):
    return {"nome": nome, "colunas": [col("Nome", "texto")] + [col(n, t) for n, t in COLUNAS_SISTEMA],
            "etapas": [{"nome": "Novo", "tipo": "aberta", "cor": "#60a5fa"},
                       {"nome": "Ganho", "tipo": "ganha", "cor": "#22c55e"},
                       {"nome": "Perdido", "tipo": "perdida", "cor": "#ef4444"}],
            "registros": []}


def novo_id():
    return uuid.uuid4().hex[:10]


def cobranca(v):
    return "mensal" if str(v or "").strip().lower().startswith("mens") else "unico"


def brl(v):
    s = f"{v:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")
    return "R$ " + (s[:-3] if s.endswith(",00") else s)


def resumo_servicos(itens):
    """Texto da célula no Excel: "Landing Page R$ 500; Tráfego Pago R$ 800/mês". Os itens ficam na aba _Servicos."""
    partes = []
    for it in itens:
        v = it.get("valor")
        preco = brl(v) + ("/mês" if it.get("cobranca") == "mensal" else "") if isinstance(v, (int, float)) else "(a definir)"
        partes.append(f'{it["servico"]} {preco}')
    return "; ".join(partes)


def mtime():
    return str(ARQ.stat().st_mtime_ns) if ARQ.exists() else None


# ---------- conversão de valores Excel <-> JSON ----------

def ler_valor(v, tipo):
    """Valor da célula -> valor do app (datas viram 'aaaa-mm-dd')."""
    if v is None or (isinstance(v, str) and not v.strip()):
        return None
    if tipo in ("dias_na_etapa", "servicos"):
        return None  # calculado / vem da aba _Servicos
    if tipo in TIPOS_DATA:
        if isinstance(v, dt.datetime):
            return v.date().isoformat()
        if isinstance(v, dt.date):
            return v.isoformat()
        if isinstance(v, (int, float)) and not isinstance(v, bool):  # número de série do Excel
            return (dt.date(1899, 12, 30) + dt.timedelta(days=int(v))).isoformat()
        m = re.fullmatch(r"\s*(\d{1,2})/(\d{1,2})/(\d{2}|\d{4})\s*", str(v))
        if m:
            d, mes, a = (int(x) for x in m.groups())
            try:
                return dt.date(a + 2000 if a < 100 else a, mes, d).isoformat()
            except ValueError:
                pass
        return str(v).strip()
    if tipo == "caixa":
        if isinstance(v, bool):
            return v
        return str(v).strip().lower() in ("sim", "s", "x", "true", "verdadeiro", "1", "ok")
    if tipo == "numero":
        if isinstance(v, (int, float)) and not isinstance(v, bool):
            return v
        try:
            return float(str(v).replace(".", "").replace(",", ".")) if "," in str(v) else float(v)
        except ValueError:
            return str(v).strip()
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    if isinstance(v, (dt.datetime, dt.date)):
        return v.strftime("%d/%m/%Y")
    return str(v).strip()


def limpo(s):
    return ILLEGAL_CHARACTERS_RE.sub("", s) if isinstance(s, str) else s


def link(v):
    v = v.strip()
    if re.match(r"(?i)https?://", v):
        return v
    return None if re.match(r"(?i)[a-z][a-z0-9+.-]*:", v) else "https://" + v


def gravar_valor(cell, v, tipo):
    if v is None or v == "":
        return
    v = limpo(v)
    if tipo in TIPOS_DATA and isinstance(v, str) and re.fullmatch(r"\d{4}-\d{2}-\d{2}", v):
        cell.value = dt.datetime.strptime(v, "%Y-%m-%d")
        cell.number_format = "DD/MM/YYYY"
        return
    cell.value = v
    if isinstance(v, str) and v.startswith("="):
        cell.data_type = "s"  # texto, não fórmula
    if tipo == "url" and isinstance(v, str) and link(v):
        cell.hyperlink = link(v)
        cell.style = "Hyperlink"


# ---------- leitura ----------

def ler_config(wb):
    cfg = {}
    if ABA_CONFIG not in wb.sheetnames:
        return cfg
    linhas = [list(l) + [None] * 12 for l in wb[ABA_CONFIG].iter_rows(min_row=2, values_only=True)]

    def ordem(l):
        try:
            return float(l[5])
        except (TypeError, ValueError):
            return 9999

    def sim(v):
        return str(v).strip().lower() in ("true", "verdadeiro", "sim", "1")

    def numero(v):
        try:
            return float(v) if float(v) != int(float(v)) else int(float(v))
        except (TypeError, ValueError):
            return None

    for lista, item, nome, tipo, opcoes, _, visivel, cor, oferta, producao, limite, padrao in (l[:12] for l in sorted(linhas, key=ordem)):
        if not lista or not nome:
            continue
        c = cfg.setdefault(str(lista).strip(), {"colunas": [], "etapas": [], "servicos": []})
        nome, tipo, item = str(nome).strip(), str(tipo or "").strip().lower(), str(item or "").strip().lower()
        if item == "papel":  # papel fixo da lista (contabilidades, clientes): vale mesmo se a lista for renomeada
            c["papel"] = nome
        elif item in ("servico", "serviço"):  # catálogo da coluna de serviços; Padrão = preço sugerido
            if nome not in [s["nome"] for s in c["servicos"]]:
                c["servicos"].append({"nome": nome, "cobranca": cobranca(tipo), "preco": numero(padrao)})
        elif item == "etapa":
            if nome not in [e["nome"] for e in c["etapas"]]:
                c["etapas"].append({"nome": nome, "tipo": tipo if tipo in TIPOS_ETAPA else "aberta",
                                    "cor": str(cor or "#94a3b8"), "oferta": sim(oferta), "producao": sim(producao),
                                    "limite": numero(limite)})
        elif nome not in [x["nome"] for x in c["colunas"]]:
            ops = [o.strip() for o in str(opcoes or "").split(";") if o.strip()]
            vis = str(visivel).strip().lower() not in ("false", "falso", "não", "nao", "0")
            c["colunas"].append(col(nome, tipo if tipo in TIPOS_COLUNA + TIPOS_SISTEMA else "texto", ops, vis))
            if padrao not in (None, ""):
                c["colunas"][-1]["padrao"] = padrao
    return cfg


def completar(lista):
    """Garante colunas de sistema, etapas e opções para valores digitados à mão no Excel."""
    mudou = False
    tipos = [c["tipo"] for c in lista["colunas"]]
    for nome, tipo in COLUNAS_SISTEMA:
        if tipo not in tipos:
            while nome in [c["nome"] for c in lista["colunas"]]:
                nome += " (2)"
            lista["colunas"].append(col(nome, tipo))
            mudou = True
    if not lista["etapas"]:
        lista["etapas"] = nova_lista("")["etapas"]
        mudou = True
    cs = next((c for c in lista["colunas"] if c["tipo"] == "servicos"), None)
    if cs:  # serviço digitado à mão na aba _Servicos entra no catálogo
        cs.setdefault("servicos", [])
        for r in lista["registros"]:
            for it in r.get(cs["nome"]) or []:
                if it["servico"] not in [s["nome"] for s in cs["servicos"]]:
                    cs["servicos"].append({"nome": it["servico"], "cobranca": it["cobranca"], "preco": None})
                    mudou = True
    nome_etapa = next(c["nome"] for c in lista["colunas"] if c["tipo"] == "etapa")
    for r in lista["registros"]:
        v = r.get(nome_etapa)
        if v and v not in [e["nome"] for e in lista["etapas"]]:
            abertas = sum(1 for e in lista["etapas"] if e["tipo"] == "aberta")
            lista["etapas"].insert(abertas, {"nome": v, "tipo": "aberta", "cor": "#94a3b8"})
            mudou = True
        for c in lista["colunas"]:
            v = r.get(c["nome"])
            if c["tipo"] == "selecao" and v and v not in c["opcoes"]:
                c["opcoes"].append(v)
                mudou = True
    return mudou


def ler_contatos(wb):
    """Contatos da aba _Contatos; devolve (contatos, mudou) — mudou quando algum _id precisou ser criado."""
    contatos, ids, mudou = [], set(), False
    if ABA_CONTATOS not in wb.sheetnames:
        return contatos, mudou
    for l in wb[ABA_CONTATOS].iter_rows(min_row=2, values_only=True):
        l = list(l)[:8] + [None] * (8 - len(l))
        tipo, ref = (str(x).strip() if x is not None else "" for x in l[:2])
        if not tipo or not ref:
            continue
        k = {"tipo": tipo, "ref": ref}
        for f, v in zip(CAMPOS_CONTATO, l[2:7]):
            v = ler_valor(v, "texto")
            if v is not None:
                k[f] = str(v)
        rid = str(l[7]).strip() if l[7] is not None else ""
        if not rid or rid in ids:
            rid, mudou = novo_id(), True
        ids.add(rid)
        k["_id"] = rid
        contatos.append(k)
    return contatos, mudou


def ler():
    if not ARQ.exists():
        return {"listas": [nova_lista("Leads")], "contatos": [], "mtime": None, "precisaSalvar": True, "versaoApi": VERSAO_API}
    versao = mtime()
    wb = openpyxl.load_workbook(ARQ)
    cfg = ler_config(wb)
    hist = {}
    if ABA_HIST in wb.sheetnames:
        for lista, rid, data, de, para in (list(l)[:5] + [None] * (5 - len(l)) for l in wb[ABA_HIST].iter_rows(min_row=2, values_only=True)):
            if lista and rid:
                hist.setdefault(str(lista), []).append({"id": str(rid), "data": ler_valor(data, "data"), "de": de or "", "para": para or ""})
    itens = {}
    if ABA_ITENS in wb.sheetnames:
        for lista, rid, servico, cob, valor in (list(l)[:5] + [None] * (5 - len(l)) for l in wb[ABA_ITENS].iter_rows(min_row=2, values_only=True)):
            if lista and rid and servico is not None and str(servico).strip():
                it = {"servico": str(servico).strip(), "cobranca": cobranca(cob)}
                v = ler_valor(valor, "numero")
                if isinstance(v, (int, float)):
                    it["valor"] = v
                itens.setdefault((str(lista), str(rid).strip()), []).append(it)
    contatos, mudou = ler_contatos(wb)
    listas = []
    por_nome = {n.casefold(): t for n, t in COLUNAS_SISTEMA}
    for ws in wb.worksheets:
        if ws.title.startswith("_"):
            continue
        c = cfg.get(ws.title)
        if c is None:
            c, mudou = {"colunas": [], "etapas": [], "servicos": []}, True
        cab = [str(x.value).strip() if x.value is not None else "" for x in ws[1]]
        for h in cab:  # coluna criada à mão no Excel
            if h and h != "_id" and h not in [x["nome"] for x in c["colunas"]]:
                tipo = por_nome.get(h.casefold(), "texto")
                if tipo in [x["tipo"] for x in c["colunas"]]:
                    tipo = "texto"
                c["colunas"].append(col(h, tipo))
                mudou = True
        idx = {h: i for i, h in enumerate(cab) if h}
        cs = next((x for x in c["colunas"] if x["tipo"] == "servicos"), None)
        if cs:
            cs["servicos"] = c.get("servicos", [])
        registros, ids = [], set()
        for linha in ws.iter_rows(min_row=2, values_only=True):
            r = {}
            for x in c["colunas"]:
                i = idx.get(x["nome"])
                v = ler_valor(linha[i], x["tipo"]) if i is not None and i < len(linha) else None
                if v is not None:
                    r[x["nome"]] = v
            if not r:
                continue
            i = idx.get("_id")
            rid = str(linha[i]).strip() if i is not None and i < len(linha) and linha[i] is not None else ""
            if not rid or rid in ids:  # vazio ou linha copiada no Excel
                rid, mudou = novo_id(), True
            ids.add(rid)
            r["_id"] = rid
            if cs and (ws.title, rid) in itens:
                r[cs["nome"]] = itens[(ws.title, rid)]
            registros.append(r)
        lista = {"nome": ws.title, "colunas": c["colunas"], "etapas": c["etapas"], "registros": registros,
                 "historico": hist.get(ws.title, [])}
        if c.get("papel"):
            lista["papel"] = c["papel"]
        mudou = completar(lista) or mudou
        listas.append(lista)
    if not listas:
        listas, mudou = [nova_lista("Leads")], True
    return {"listas": listas, "contatos": contatos, "mtime": versao, "precisaSalvar": mudou, "versaoApi": VERSAO_API}


# ---------- gravação ----------

def validar(estado):
    listas = estado.get("listas") or []
    if not listas:
        raise ValueError("É preciso ter pelo menos uma lista.")
    vistos, papeis = set(), set()
    for l in listas:
        n = l.get("nome", "")
        if not n or len(n) > 31 or re.search(r"[\[\]:*?/\\]", n) or n.startswith(("_", "'")) or n.endswith("'"):
            raise ValueError(f'Nome de lista inválido: "{n}" (até 31 caracteres, sem [ ] : * ? / \\ e sem começar com _).')
        if n.casefold() in vistos:
            raise ValueError(f'Lista repetida: "{n}".')
        vistos.add(n.casefold())
        if l.get("papel") and l["papel"] in papeis:
            raise ValueError(f'Duas listas com o papel "{l["papel"]}".')
        if l.get("papel"):
            papeis.add(l["papel"])
        nomes = [c["nome"].casefold() for c in l["colunas"]]
        if any(not c["nome"] or c["nome"].startswith("_") for c in l["colunas"]) or len(set(nomes)) != len(nomes):
            raise ValueError(f'Lista "{n}": nomes de coluna vazios, repetidos ou começando com _.')
        tipos = [c["tipo"] for c in l["colunas"]]
        if (any(t not in TIPOS_COLUNA + TIPOS_SISTEMA for t in tipos) or any(tipos.count(t) != 1 for t in TIPOS_SISTEMA)
                or tipos.count("servicos") > 1 or tipos.count("vinculo") > 1):
            raise ValueError(f'Lista "{n}": tipos de coluna inválidos.')
        etapas = [e["nome"].casefold() for e in l["etapas"]]
        if not etapas or len(set(etapas)) != len(etapas) or any(e["tipo"] not in TIPOS_ETAPA for e in l["etapas"]):
            raise ValueError(f'Lista "{n}": etapas vazias, repetidas ou com tipo inválido.')


def formatar(ws, colunas):
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    cinza = PatternFill("solid", fgColor="E5E7EB")
    for cell in ws[1]:
        cell.font, cell.fill = Font(bold=True), cinza
    for i, c in enumerate(colunas + [col("_id", "texto")], start=1):
        letra = get_column_letter(i)
        if c["tipo"] in TIPOS_DATA or c["tipo"] == "dias_na_etapa":
            largura = max(12, len(c["nome"]) + 2)
        else:
            maior = max((len(str(x.value)) for x in ws[letra][:200] if x.value is not None), default=0)
            largura = min(50, max(10, maior + 2))
        ws.column_dimensions[letra].width = largura
    for cell in ws[get_column_letter(len(colunas) + 1)]:
        cell.font = Font(color="9CA3AF", bold=cell.row == 1)


def fazer_backup():
    global ultimo_backup
    if not ARQ.exists() or time.time() - ultimo_backup < INTERVALO_BACKUP:
        return
    BACKUPS.mkdir(exist_ok=True)
    shutil.copy2(ARQ, BACKUPS / f"crm_{time.strftime('%Y-%m-%d_%H%M%S')}.xlsx")
    ultimo_backup = time.time()
    for velho in sorted(BACKUPS.glob("crm_*.xlsx"))[:-MAX_BACKUPS]:
        velho.unlink()


def gravar(estado):
    validar(estado)
    contatos = estado.get("contatos")
    if not isinstance(contatos, list):  # um CRM de versão antiga não manda os contatos: mantém os do arquivo
        contatos = ler_contatos(openpyxl.load_workbook(ARQ))[0] if ARQ.exists() else []
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    config, itens = [], []
    for l in estado["listas"]:
        ws = wb.create_sheet(l["nome"])
        cols = l["colunas"]
        if l.get("papel"):
            config.append([l["nome"], "papel", l["papel"], None, None, 0, None, None, None, None, None, None])
        ws.append([limpo(c["nome"]) for c in cols] + ["_id"])
        desde = next(get_column_letter(i) for i, c in enumerate(cols, 1) if c["tipo"] == "etapa_desde")
        for n, r in enumerate(l["registros"], start=2):
            for i, c in enumerate(cols, start=1):
                cell = ws.cell(n, i)
                if c["tipo"] == "dias_na_etapa":
                    cell.value = f'=IF({desde}{n}="","",TODAY()-{desde}{n})'
                    cell.number_format = "0"
                elif c["tipo"] == "servicos":
                    lista_itens = [it for it in r.get(c["nome"]) or [] if isinstance(it, dict) and it.get("servico")]
                    gravar_valor(cell, resumo_servicos(lista_itens), "texto")
                    for it in lista_itens:
                        v = it.get("valor")
                        itens.append([l["nome"], str(r["_id"]), it["servico"], COBRANCAS[cobranca(it.get("cobranca"))],
                                      v if isinstance(v, (int, float)) and not isinstance(v, bool) else None])
                else:
                    gravar_valor(cell, r.get(c["nome"]), c["tipo"])
            ws.cell(n, len(cols) + 1, limpo(str(r["_id"])))
        formatar(ws, cols)
        for i, c in enumerate(cols, 1):
            config.append([l["nome"], "coluna", c["nome"], c["tipo"], "; ".join(c["opcoes"]), i, bool(c["visivel"]),
                           None, None, None, None, c.get("padrao")])
            for j, s in enumerate(c.get("servicos") or [] if c["tipo"] == "servicos" else [], 1):
                config.append([l["nome"], "servico", s["nome"], cobranca(s.get("cobranca")), None, j, None,
                               None, None, None, None, s.get("preco")])
        for i, e in enumerate(l["etapas"], 1):
            config.append([l["nome"], "etapa", e["nome"], e["tipo"], None, i, None, e["cor"], bool(e.get("oferta")),
                           bool(e.get("producao")), e.get("limite"), None])
    wh = wb.create_sheet(ABA_HIST)
    wh.append(CAB_HIST)
    for l in estado["listas"]:
        for h in l.get("historico") or []:
            n = wh.max_row + 1
            wh.append([limpo(l["nome"]), limpo(str(h.get("id", ""))), None, limpo(h.get("de")), limpo(h.get("para"))])
            gravar_valor(wh.cell(n, 3), h.get("data"), "data")
    formatar(wh, [col(n, "texto") for n in CAB_HIST])
    wi = wb.create_sheet(ABA_ITENS)
    wi.append(CAB_ITENS)
    for linha in itens:
        wi.append([limpo(x) for x in linha])
    formatar(wi, [col(n, "texto") for n in CAB_ITENS])
    wk = wb.create_sheet(ABA_CONTATOS)
    wk.append(CAB_CONTATOS)
    for k in contatos:
        if isinstance(k, dict) and k.get("tipo") and k.get("ref"):
            wk.append([limpo(str(k["tipo"])), limpo(str(k["ref"]))] + [limpo(str(k[f])) if k.get(f) is not None else None for f in CAMPOS_CONTATO]
                      + [limpo(str(k.get("_id") or novo_id()))])
    for linha in wk.iter_rows(min_row=2):
        for cell in linha:
            cell.number_format = "@"
    formatar(wk, [col(n, "texto") for n in CAB_CONTATOS])
    wc = wb.create_sheet(ABA_CONFIG)
    wc.append(CAB_CONFIG)
    for linha in config:
        wc.append([limpo(x) for x in linha])
    formatar(wc, [col(n, "texto") for n in CAB_CONFIG])

    tmp = ARQ.with_name("crm.tmp.xlsx")
    wb.save(tmp)
    try:
        fazer_backup()
        os.replace(tmp, ARQ)  # troca atômica: ou fica o arquivo antigo inteiro, ou o novo inteiro
    except PermissionError:
        tmp.unlink(missing_ok=True)
        raise ArquivoTravado()


# ---------- HTTP ----------

class Handler(BaseHTTPRequestHandler):
    def log_message(self, *args):
        pass

    def responder(self, codigo, obj=None, corpo=None, tipo="application/json; charset=utf-8"):
        dados = corpo if corpo is not None else json.dumps(obj, ensure_ascii=False).encode("utf-8")
        self.send_response(codigo)
        self.send_header("Content-Type", tipo)
        self.send_header("Content-Length", str(len(dados)))
        self.send_header("Cache-Control", "no-store")
        self.end_headers()
        self.wfile.write(dados)

    def origem_ok(self):
        # Bloqueia requisições vindas de outros sites abertos no navegador.
        locais = (f"127.0.0.1:{PORTA}", f"localhost:{PORTA}")
        origem = self.headers.get("Origin")
        return self.headers.get("Host") in locais and (origem is None or origem in ["http://" + h for h in locais])

    def do_GET(self):
        if not self.origem_ok():
            return self.responder(403, {"erro": "Acesso negado."})
        if self.path in ("/", "/index.html"):
            self.responder(200, corpo=(PASTA / "index.html").read_bytes(), tipo="text/html; charset=utf-8")
        elif self.path == "/api/dados":
            try:
                self.responder(200, ler())
            except Exception as e:
                self.responder(500, {"erro": f"Não consegui ler o crm.xlsx: {e}"})
        else:
            self.responder(404, {"erro": "Não encontrado."})

    def do_POST(self):
        if not self.origem_ok() or self.headers.get("Content-Type", "").split(";")[0] != "application/json":
            return self.responder(403, {"erro": "Acesso negado."})
        if self.path != "/api/salvar":
            return self.responder(404, {"erro": "Não encontrado."})
        try:
            estado = json.loads(self.rfile.read(int(self.headers.get("Content-Length", 0))))
        except ValueError:
            return self.responder(400, {"erro": "Dados inválidos."})
        if ARQ.exists() and not estado.get("forcar") and estado.get("mtime") != mtime():
            return self.responder(409, {"erro": "O crm.xlsx foi alterado fora do app (no Excel?) depois da última leitura."})
        try:
            gravar(estado)
        except ValueError as e:
            return self.responder(400, {"erro": str(e)})
        except ArquivoTravado:
            return self.responder(423, {"erro": "Não consegui gravar: o crm.xlsx parece estar aberto no Excel. Feche o arquivo."})
        except Exception as e:
            return self.responder(500, {"erro": f"Erro ao gravar: {e}"})
        self.responder(200, {"mtime": mtime()})


if __name__ == "__main__":
    try:
        servidor = HTTPServer(("127.0.0.1", PORTA), Handler)
    except OSError:
        print("O CRM já está rodando. Abrindo o navegador...")
        webbrowser.open(URL)
        sys.exit(0)
    print(f"CRM rodando em {URL}\nArquivo: {ARQ}\nFeche esta janela para encerrar.")
    if "--sem-navegador" not in sys.argv:
        webbrowser.open(URL)
    servidor.serve_forever()
